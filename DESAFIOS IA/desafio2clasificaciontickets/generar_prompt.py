"""Arma el prompt completo para ChatGPT a partir de tickets.xlsx.
Uso: python generar_prompt.py tickets.xlsx prompt.txt
Requiere: pip install openpyxl
"""
import sys
from openpyxl import load_workbook

INSTRUCCIONES = (
    "Eres un analista de soporte de TI. Clasifica cada ticket en una sola categoría: "
    "Hardware, Software, Red, Acceso/Cuentas, Seguridad o Consulta general.\n"
    "Responde solo con una tabla de dos columnas: ID | Tipo de incidencia.\n"
    "Si un ticket es ambiguo, elige la categoría más probable.\n"
)

def leer_tickets(ruta):
    ws = load_workbook(ruta).active
    tickets = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        id_, _fecha, _usuario, mensaje = fila[:4]
        if id_ is None or not mensaje:
            continue
        tickets.append((int(id_), str(mensaje).strip()))
    return tickets

def armar_prompt(tickets):
    lineas = [f"{id_} | {mensaje}" for id_, mensaje in tickets]
    return INSTRUCCIONES + "\nTickets:\n" + "\n".join(lineas) + "\n"

if __name__ == "__main__":
    entrada = sys.argv[1] if len(sys.argv) > 1 else "tickets.xlsx"
    salida = sys.argv[2] if len(sys.argv) > 2 else "prompt.txt"
    tickets = leer_tickets(entrada)
    with open(salida, "w", encoding="utf-8") as f:
        f.write(armar_prompt(tickets))
    print(f"Prompt con {len(tickets)} tickets guardado en {salida}")
