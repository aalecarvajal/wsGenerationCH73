"""Escribe en el Excel la clasificación que devolvió ChatGPT (columna E).
Uso: python llenar_clasificacion.py tickets.xlsx respuesta.txt [salida.xlsx]
  - respuesta.txt: pega ahí la tabla "ID | Tipo de incidencia" de ChatGPT.
  - Si no das salida.xlsx, se sobrescribe tickets.xlsx.
Requiere: pip install openpyxl
"""
import re
import sys
import unicodedata
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

CATEGORIAS = ["Hardware", "Software", "Red", "Acceso/Cuentas", "Seguridad", "Consulta general"]
AMARILLO = PatternFill("solid", start_color="FFFF00")

def normalizar(texto):
    t = unicodedata.normalize("NFD", texto)
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"[\s*_`]+", " ", t).strip().lower()

CANON = {normalizar(c): c for c in CATEGORIAS}

def parsear_respuesta(texto):
    """Devuelve {id: categoria_cruda} leyendo líneas tipo '1001 | Software' o '| 1001 | Software |'."""
    resultado = {}
    for linea in texto.splitlines():
        partes = [p.strip() for p in linea.strip().strip("|").split("|")]
        if len(partes) < 2:
            continue
        m = re.fullmatch(r"\*{0,2}#?(\d+)\*{0,2}", partes[0])
        if m:
            resultado[int(m.group(1))] = partes[1]
    return resultado

def llenar(ruta_xlsx, ruta_respuesta, ruta_salida):
    with open(ruta_respuesta, encoding="utf-8") as f:
        respuestas = parsear_respuesta(f.read())
    wb = load_workbook(ruta_xlsx)
    ws = wb.active
    sin_respuesta, invalidas, ids_excel = [], [], set()
    for fila in range(2, ws.max_row + 1):
        id_ = ws.cell(fila, 1).value
        if id_ is None:
            continue
        id_ = int(id_)
        ids_excel.add(id_)
        if id_ not in respuestas:
            sin_respuesta.append(id_)
            continue
        cruda = respuestas[id_]
        celda = ws.cell(fila, 5)
        canonica = CANON.get(normalizar(cruda))
        if canonica:
            celda.value = canonica
        else:
            celda.value = cruda
            celda.fill = AMARILLO
            invalidas.append((id_, cruda))
    wb.save(ruta_salida)
    extras = sorted(set(respuestas) - ids_excel)
    print(f"{len(ids_excel) - len(sin_respuesta)} de {len(ids_excel)} tickets clasificados -> {ruta_salida}")
    if sin_respuesta:
        print("Sin respuesta en el texto:", sin_respuesta)
    if invalidas:
        print("Categoría no reconocida (marcada en amarillo):", invalidas)
    if extras:
        print("IDs de la respuesta que no existen en el Excel:", extras)

if __name__ == "__main__":
    xlsx = sys.argv[1] if len(sys.argv) > 1 else "tickets.xlsx"
    resp = sys.argv[2] if len(sys.argv) > 2 else "respuesta.txt"
    salida = sys.argv[3] if len(sys.argv) > 3 else xlsx
    llenar(xlsx, resp, salida)
